from __future__ import annotations

import argparse
import re
from pathlib import Path

from common import OUT_DIR, ensure_dirs


ACRONYMS = ["BTMM", "DRQ", "RWS", "MSX", "RDX", "SAC-D", "MARL", "XRL"]
FULL_FORMS = {
    "BTMM": "binary tree merged monte carlo",
    "DRQ": "decomposed reward q",
    "RWS": "roulette wheel selection",
    "MSX": "minimal sufficient explanation",
    "RDX": "reward decomposition explanation",
    "SAC-D": "soft actor-critic discrete",
    "MARL": "multi-agent reinforcement learning",
    "XRL": "explainable reinforcement learning",
}


def docx_text(path: Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(path)
    chunks: list[str] = []
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            chunks.append(Paragraph(child, doc).text)
        elif child.tag.endswith("}tbl"):
            table = Table(child, doc)
            for row in table.rows:
                chunks.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(chunks)


def citation_numbers(text: str) -> set[int]:
    found: set[int] = set()
    for body in re.findall(r"\[([0-9,\s\-–]+)\]", text):
        if re.search(r"\b0\b", body):
            continue
        for piece in re.split(r",", body):
            piece = piece.strip()
            range_match = re.match(r"^(\d+)\s*[-–]\s*(\d+)$", piece)
            if range_match:
                start, end = map(int, range_match.groups())
                found.update(range(start, end + 1))
            elif piece.isdigit():
                found.add(int(piece))
    return found


def reference_numbers(text: str) -> set[int]:
    explicit_refs: set[int] = set()
    inferred_entries: list[str] = []
    pending = ""
    in_refs = False
    for line in text.splitlines():
        line = line.strip()
        if re.match(r"^\s*references\s*$", line, re.I):
            in_refs = True
            continue
        if not in_refs:
            continue
        match = re.match(r"^\s*\[?(\d+)\]?[.)]?\s+", line)
        if match:
            explicit_refs.add(int(match.group(1)))
            continue
        if not line or re.match(r"^URL\b", line, re.I):
            continue
        if not pending:
            pending = line
        elif pending.rstrip().endswith(","):
            pending += " " + line
        else:
            inferred_entries.append(pending)
            pending = line
    if pending:
        inferred_entries.append(pending)
    if explicit_refs:
        return explicit_refs
    return set(range(1, len(inferred_entries) + 1))


def caption_numbers(text: str, kind: str) -> set[int]:
    return {int(value) for value in re.findall(rf"\b{kind}\s+(\d+)\s*[:.]", text, re.I)}


def mentioned_numbers(text: str, kind: str) -> set[int]:
    return {int(value) for value in re.findall(rf"\b{kind}\s+(\d+)\b", text, re.I)}


def acronym_checks(text: str) -> list[tuple[str, str, str]]:
    lower = text.lower()
    rows: list[tuple[str, str, str]] = []
    for acronym in ACRONYMS:
        match = re.search(rf"\b{re.escape(acronym)}\b", text)
        if not match:
            rows.append((acronym, "not found", "OK"))
            continue
        window = lower[max(0, match.start() - 140) : match.end() + 40]
        full = FULL_FORMS[acronym]
        status = "OK" if full in window else "WARN first occurrence may lack full form"
        rows.append((acronym, f"char {match.start()}", status))
    return rows


def markdown_table(headers: list[str], rows: list[tuple[object, ...]]) -> str:
    out = ["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"]
    for row in rows:
        out.append("| " + " | ".join(str(item) for item in row) + " |")
    return "\n".join(out)


def main() -> None:
    parser = argparse.ArgumentParser(description="Check manuscript citations, captions, placeholders, and acronym first uses.")
    parser.add_argument("docx")
    args = parser.parse_args()
    ensure_dirs()
    path = Path(args.docx)
    text = docx_text(path)
    cited = citation_numbers(text)
    refs = reference_numbers(text)
    missing_refs = sorted(cited - refs)
    uncited_refs = sorted(refs - cited)
    valid_range = sorted(num for num in cited if refs and (num < min(refs) or num > max(refs)))
    figure_problems = sorted(mentioned_numbers(text, "Figure") - caption_numbers(text, "Figure"))
    table_problems = sorted(mentioned_numbers(text, "Table") - caption_numbers(text, "Table"))

    sections = [
        "# Manuscript Check",
        markdown_table(
            ["Check", "Result"],
            [
                ("Cited numbers missing from references", ", ".join(map(str, missing_refs)) or "none"),
                ("References not cited in text", ", ".join(map(str, uncited_refs)) or "none"),
                ("Citation numbers outside reference range", ", ".join(map(str, valid_range)) or "none"),
                ("Mentioned figures without captions", ", ".join(map(str, figure_problems)) or "none"),
                ("Mentioned tables without captions", ", ".join(map(str, table_problems)) or "none"),
                ("Remaining __ placeholders", text.count("__")),
            ],
        ),
        "## Acronym First Uses",
        markdown_table(["Acronym", "First occurrence", "Status"], acronym_checks(text)),
    ]
    report = "\n\n".join(sections) + "\n"
    (OUT_DIR / "manuscript_check.md").write_text(report, encoding="utf-8")
    print(report)
    print(f"Wrote {OUT_DIR / 'manuscript_check.md'}")


if __name__ == "__main__":
    main()
