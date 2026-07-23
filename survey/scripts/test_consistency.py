from __future__ import annotations

import os
from pathlib import Path

from common import EXTRACTION, OUT_DIR, read_csv, read_json


def manuscript_text() -> str:
    path_value = os.environ.get("SURVEY_MANUSCRIPT_DOCX")
    assert path_value, "Set SURVEY_MANUSCRIPT_DOCX to the manuscript .docx path"
    from docx import Document

    document = Document(Path(path_value))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def test_screening_arithmetic() -> None:
    summary = read_json(OUT_DIR / "screening_summary.json")
    exclusions = summary["exclusions"]
    assert sum(exclusions.values()) == summary["n_title_abstract_pass"] - summary["n_fulltext_pass"], (
        "E1+E2+E3+E4 must equal title/abstract pass minus full-text pass"
    )


def test_extraction_count_matches_fulltext_pass() -> None:
    summary = read_json(OUT_DIR / "screening_summary.json")
    rows = read_csv(EXTRACTION)
    assert len(rows) == summary["n_fulltext_pass"], "extraction.csv row count must equal n_fulltext_pass"


def test_manuscript_has_no_placeholders() -> None:
    text = manuscript_text()
    assert "__" not in text, "manuscript still contains __ placeholders"


def test_key_screening_numbers_appear_in_manuscript() -> None:
    summary = read_json(OUT_DIR / "screening_summary.json")
    text = manuscript_text()
    expected_numbers = [
        summary["n_identified"],
        summary["n_duplicates"],
        summary["n_after_dedup"],
        summary["n_title_abstract_pass"],
        summary["n_fulltext_pass"],
        *summary["exclusions"].values(),
    ]
    for number in expected_numbers:
        assert str(number) in text, f"Expected survey count {number} to appear in manuscript"
